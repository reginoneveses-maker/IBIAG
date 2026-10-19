import io
import os
import re
import unittest
import zipfile
from unittest.mock import patch
from fastapi import UploadFile, HTTPException
from openpyxl import Workbook
import tabular_imports as tables


def workbook_bytes(rows, name="Sheet"):
    book = Workbook()
    book.active.title = name
    for row in rows:
        book.active.append(row)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def rewrite_package(raw, replacements=None, extras=None):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(output,"w",zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry)
            if replacements and entry.filename in replacements:
                data = replacements[entry.filename](data)
            target.writestr(entry.filename,data)
        for name,data in (extras or {}).items():
            target.writestr(name,data)
    return output.getvalue()


async def read(raw, name="data.xlsx", **kwargs):
    return await tables.read_table(UploadFile(io.BytesIO(raw),filename=name), **kwargs)


class TabularReaderTests(unittest.IsolatedAsyncioTestCase):
    async def test_two_thousand_companies_fit_without_truncation(self):
        rows=[["Company","Phone"],*[[f"Synthetic {i}","001234"] for i in range(2000)]]
        text=io.StringIO(newline="")
        import csv
        csv.writer(text).writerows(rows)
        for raw,name in ((text.getvalue().encode(),"data.csv"),(workbook_bytes(rows),"data.xlsx")):
            frame=await read(raw,name)
            self.assertEqual(frame.shape,(2000,2))
            self.assertEqual(frame.iloc[-1].tolist(),["Synthetic 1999","001234"])

    async def test_csv_quotes_multiline_utf8_and_leading_zeroes(self):
        raw='\ufeffCompany;Phone;Notes\r\n"Açúcar, Brasil";001234;"linha 1\nlinha 2; ""citação"""\r\n'.encode()
        frame=await read(raw,"data.csv")
        self.assertEqual(frame.iloc[0].tolist(),["Açúcar, Brasil","001234",'linha 1\nlinha 2; "citação"'])

    async def test_csv_limits_reject_before_row_reader_or_dataframe(self):
        cases=[("MAX_ROWS",1,b"Name\na\nb\n"),("MAX_COLUMNS",2,b"a,b,c\n1,2,3\n"),
               ("MAX_CELLS",3,b"a,b\n1,2\n"),("MAX_TEXT_CHARS",4,b"Name\na\n"),
               ("MAX_CELL_CHARS",4,b"Name\n12345\n")]
        for setting,value,raw in cases:
            with patch.object(tables,setting,value),patch.object(tables.csv,"reader",side_effect=AssertionError("reader must not run")):
                with self.assertRaises(HTTPException) as error: await read(raw,"data.csv")
                self.assertEqual(error.exception.status_code,413,(setting,error.exception.detail))

    async def test_csv_row_boundary_and_quoted_delimiters_are_not_truncated(self):
        with patch.object(tables,"MAX_ROWS",2),patch.object(tables,"MAX_COLUMNS",2):
            frame=await read(b'Name,Notes\na,"1,2,3"\nb,"4\n5"\n',"data.csv")
            self.assertEqual(frame.shape,(2,2))
            self.assertEqual(frame.iloc[1]["Notes"],"4\n5")

    async def test_sparse_table_cannot_expand_past_grid_budget_in_dataframe(self):
        rows=[["A","B","C","D"],["one"],["two"],["three"]]
        for raw,name in ((b"A,B,C,D\none\ntwo\nthree\n","data.csv"),(workbook_bytes(rows),"data.xlsx")):
            with patch.object(tables,"MAX_CELLS",8),patch.object(tables.pd,"DataFrame",side_effect=AssertionError("frame must not allocate")) as frame:
                with self.assertRaises(HTTPException) as error: await read(raw,name)
                self.assertEqual(error.exception.status_code,413)
                frame.assert_not_called()

    async def test_malformed_duplicate_headers_binary_xls_and_wrong_encoding_are_generic(self):
        for raw,name in ((b'Company,company\na,b\n',"data.csv"),(b'Company\n"private-canary',"data.csv"),
                         (b'Company\n\xff',"data.csv"),(b"private-canary","data.xlsx"),(b"private-canary","old.xls")):
            with self.assertRaises(HTTPException) as error: await read(raw,name)
            self.assertEqual(error.exception.status_code,400)
            self.assertEqual(error.exception.detail,tables.INVALID_TABLE)

    async def test_xlsx_portfolio_header_and_crm_strings_are_preserved(self):
        raw=workbook_bytes([["intro"],[],["Produto","Fornecedor","Comissão"],["Açaí","Synthetic","10%"]],"IBIAG")
        frame=await read(raw,sheet_name="IBIAG",header=2)
        self.assertEqual(frame.iloc[0].tolist(),["Açaí","Synthetic","10%"])
        raw=workbook_bytes([["Company","Phone"],["Açaí","001234"]])
        self.assertEqual((await read(raw)).iloc[0]["Phone"],"001234")

    async def test_false_large_dimensions_are_ignored_without_losing_real_rows(self):
        raw=workbook_bytes([["Company"],["A"],["B"]])
        raw=rewrite_package(raw,{"xl/worksheets/sheet1.xml":lambda data:re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:XFD1048576"',data)})
        self.assertEqual((await read(raw))["Company"].tolist(),["A","B"])

    async def test_sparse_or_understated_cell_coordinates_are_rejected_before_workbook_loader(self):
        raw=workbook_bytes([["Company"],["A"]])
        for reference in (b"XFD2",b"A1048576"):
            changed=rewrite_package(raw,{"xl/worksheets/sheet1.xml":lambda data:re.sub(rb'r="A2"',b'r="'+reference+b'"',re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:A1"',data))})
            with patch.object(tables,"load_workbook",side_effect=AssertionError("loader must not run")):
                with self.assertRaises(HTTPException) as error: await read(changed)
                self.assertEqual(error.exception.status_code,413)

    async def test_xlsx_structural_budgets_precede_workbook_loading(self):
        raw=workbook_bytes([["Company","Phone"],["A","123"],["B","456"]])
        for setting,value in (("MAX_ROWS",1),("MAX_COLUMNS",1),("MAX_CELLS",3),("MAX_TEXT_CHARS",10),
                              ("MAX_XML_PART_BYTES",100),("MAX_XML_ELEMENTS",5),("MAX_XML_DEPTH",2),
                              ("MAX_STYLE_ELEMENTS",1),("MAX_METADATA_ELEMENTS",1),("MAX_SHEETS",0)):
            with patch.object(tables,setting,value),patch.object(tables,"load_workbook",side_effect=AssertionError("loader must not run")):
                with self.assertRaises(HTTPException) as error: await read(raw)
                self.assertEqual(error.exception.status_code,413,(setting,error.exception.detail))

    async def test_shared_strings_and_rich_text_are_bounded_before_openpyxl(self):
        raw=workbook_bytes([["A"],["B"]])
        shared=b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>A</t></si><si><t>B</t></si></sst>'
        changed=rewrite_package(raw,extras={"xl/sharedStrings.xml":shared})
        with patch.object(tables,"MAX_SHARED_STRINGS",1),patch.object(tables,"load_workbook",side_effect=AssertionError("loader must not run")):
            with self.assertRaises(HTTPException) as error: await read(changed)
            self.assertEqual(error.exception.status_code,413)
        shared=b'<sst><si><r><t>123</t></r><r><t>456</t></r></si></sst>'
        changed=rewrite_package(raw,extras={"xl/sharedStrings.xml":shared})
        with patch.object(tables,"MAX_CELL_CHARS",4),patch.object(tables,"load_workbook",side_effect=AssertionError("loader must not run")):
            with self.assertRaises(HTTPException) as error: await read(changed)
            self.assertEqual(error.exception.status_code,413)

    async def test_package_dtd_and_path_ambiguities_are_rejected_without_external_fetch(self):
        raw=workbook_bytes([["Company"],["A"]])
        cases=[{"unused.xml":b'<!DOCTYPE x SYSTEM "https://must-not-fetch.invalid/private.dtd"><x/>'},
               {"../private.xml":b"<x/>"}]
        with patch("urllib.request.urlopen",side_effect=AssertionError("no external fetch")) as fetch,patch.object(tables,"load_workbook",side_effect=AssertionError("loader must not run")) as loader:
            for extras in cases:
                with self.assertRaises(HTTPException) as error: await read(rewrite_package(raw,extras=extras))
                self.assertEqual(error.exception.status_code,400)
                self.assertEqual(error.exception.detail,tables.INVALID_TABLE)
            fetch.assert_not_called()
            loader.assert_not_called()
